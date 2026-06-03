import sys
import unittest
from io import BytesIO
from pathlib import Path


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from core.task_envelope import TaskEnvelope
from perception.document_processor import (
    DOCX_MIME_TYPE,
    DocumentProcessingError,
    DocumentProcessor,
)


class FakeParagraph:
    def __init__(self, text):
        self.text = text


class FakeCell:
    def __init__(self, text):
        self.text = text


class FakeRow:
    def __init__(self, cells):
        self.cells = cells


class FakeTable:
    def __init__(self, rows):
        self.rows = rows


class FakeDocument:
    paragraphs = [
        FakeParagraph("Titulo"),
        FakeParagraph(""),
        FakeParagraph("Contenido principal"),
    ]
    tables = [
        FakeTable(
            [
                FakeRow([FakeCell("A1"), FakeCell("B1")]),
                FakeRow([FakeCell("A2"), FakeCell("")]),
            ]
        )
    ]


class DocumentProcessorTests(unittest.TestCase):
    def test_docx_to_text_envelope_uses_embedded_payload(self):
        payload = b"fake-docx"
        source = TaskEnvelope.from_file_bytes(
            payload,
            filename="informe.docx",
            mime_type=DOCX_MIME_TYPE,
            declared_intent="extract_text",
            created_at="2026-05-31T12:00:00Z",
            embed_payload=True,
        )

        def fake_loader(stream):
            self.assertEqual(stream.read(), payload)
            return FakeDocument()

        processor = DocumentProcessor(docx_loader=fake_loader)
        result = processor.docx_to_text_envelope(source, created_at="2026-05-31T12:01:00Z")

        self.assertEqual(result.source_type, "text")
        self.assertEqual(result.origin, "document_processor")
        self.assertEqual(result.source_name, "informe.docx.txt")
        self.assertEqual(result.declared_intent, "plain_text_extraction")
        self.assertEqual(result.payload.value, "Titulo\nContenido principal\nA1 | B1\nA2")
        self.assertTrue(result.verify_payload_text(result.payload.value))
        self.assertEqual(result.metadata["source_task"]["task_id"], source.task_id)
        self.assertEqual(result.metadata["source_task"]["payload_sha256"], source.payload.sha256)
        self.assertEqual(result.metadata["extraction"]["paragraph_count"], 2)
        self.assertEqual(result.metadata["extraction"]["table_count"], 1)
        self.assertEqual(result.metadata["extraction"]["table_row_count"], 2)

    def test_docx_to_text_envelope_accepts_external_payload_after_hash_check(self):
        payload = b"external-docx"
        source = TaskEnvelope.from_file_bytes(
            payload,
            filename="informe.docx",
            mime_type=DOCX_MIME_TYPE,
            created_at="2026-05-31T12:00:00Z",
            embed_payload=False,
        )
        processor = DocumentProcessor(docx_loader=lambda stream: FakeDocument())

        result = processor.docx_to_text_envelope(source, payload_bytes=payload)

        self.assertEqual(result.payload.value.splitlines()[0], "Titulo")

    def test_external_ref_requires_payload_bytes(self):
        source = TaskEnvelope.from_file_bytes(
            b"external-docx",
            filename="informe.docx",
            mime_type=DOCX_MIME_TYPE,
            created_at="2026-05-31T12:00:00Z",
            embed_payload=False,
        )

        with self.assertRaises(DocumentProcessingError):
            DocumentProcessor(docx_loader=lambda stream: FakeDocument()).docx_to_text_envelope(source)

    def test_rejects_payload_bytes_that_do_not_match_source_hash(self):
        source = TaskEnvelope.from_file_bytes(
            b"original-docx",
            filename="informe.docx",
            mime_type=DOCX_MIME_TYPE,
            created_at="2026-05-31T12:00:00Z",
            embed_payload=False,
        )

        with self.assertRaises(DocumentProcessingError):
            DocumentProcessor(docx_loader=lambda stream: FakeDocument()).docx_to_text_envelope(
                source,
                payload_bytes=b"tampered-docx",
            )

    def test_rejects_non_docx_envelope(self):
        source = TaskEnvelope.from_text("hola", source_name="nota.txt")

        with self.assertRaises(DocumentProcessingError):
            DocumentProcessor(docx_loader=lambda stream: FakeDocument()).docx_to_text_envelope(source)

    def test_extracts_text_from_real_docx_bytes(self):
        from docx import Document

        buffer = BytesIO()
        document = Document()
        document.add_paragraph("Parrafo real")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Celda A"
        table.rows[0].cells[1].text = "Celda B"
        document.save(buffer)
        payload = buffer.getvalue()
        source = TaskEnvelope.from_file_bytes(
            payload,
            filename="real.docx",
            mime_type=DOCX_MIME_TYPE,
            created_at="2026-05-31T12:00:00Z",
        )

        envelope = DocumentProcessor().docx_to_text_envelope(source, payload_bytes=payload)

        self.assertEqual(envelope.payload.value, "Parrafo real\nCelda A | Celda B")
        self.assertEqual(envelope.metadata["extraction"]["paragraph_count"], 1)
        self.assertEqual(envelope.metadata["extraction"]["table_row_count"], 1)


if __name__ == "__main__":
    unittest.main()
